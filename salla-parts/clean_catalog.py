"""
تنظيف تقرير الأصناف (قطع غيار) وتحويله لملف جاهز لسلة.

الاستخدام:
    python3 clean_catalog.py input.xlsx output.xlsx

يقرأ تقرير "بيانات الأصناف" من نظام المحاسبة (يدعم أكثر من مجموعة/ماركة في نفس الملف)،
ويستخرج: نوع القطعة، الموضع، الجهة، السيارة، السنوات، الأرقام البديلة،
ويولّد: اسم محسّن، رابط، Meta Title، Meta Description، وصف، تصنيف.
الصفوف اللي فيها مشكلة تنعلّم في عمود "ملاحظات للمراجعة".
"""
import re
import sys
from collections import defaultdict

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

# ---------- القواميس (وسّعها مع باقي الملف) ----------

# الكلمة كما تُكتب في النظام -> (الماركة، الموديل بالعربي، الموديل بالإنجليزي)
MODELS = [
    (r"جيب\s*شاص|شاص", ("تويوتا", "لاندكروزر شاص", "Land Cruiser 70")),
    (r"صالون", ("تويوتا", "لاندكروزر", "Land Cruiser")),
    (r"را[فڤ]?\s*[فڤ]ور", ("تويوتا", "راف فور", "RAV4")),
    (r"افلون|افالون", ("تويوتا", "أفالون", "Avalon")),
    (r"كرولا", ("تويوتا", "كورولا", "Corolla")),
    (r"يارس", ("تويوتا", "يارس", "Yaris")),
    (r"كامري", ("تويوتا", "كامري", "Camry")),
    (r"كرسيد[اه]ء?", ("تويوتا", "كريسيدا", "Cressida")),
    (r"هايل[وو]?كس", ("تويوتا", "هايلكس", "Hilux")),
    (r"فورشنر|فورتشنر", ("تويوتا", "فورتشنر", "Fortuner")),
    (r"انوفاء?|إنوفا", ("تويوتا", "إنوفا", "Innova")),
    (r"برادو", ("تويوتا", "برادو", "Prado")),
    (r"اف\s*جي|FJ", ("تويوتا", "إف جي", "FJ Cruiser")),
    (r"RX\s*330\s*\+\s*350", ("لكزس", "RX 330/350", "RX 330/350")),
    (r"[اأ]?لكز[سز]", ("لكزس", "", "")),
    (r"النترا", ("هيونداي", "النترا", "Elantra")),
    (r"سوناتا", ("هيونداي", "سوناتا", "Sonata")),
    (r"اكسنت", ("هيونداي", "أكسنت", "Accent")),
    (r"صني", ("نيسان", "صني", "Sunny")),
    (r"باترول", ("نيسان", "باترول", "Patrol")),
    (r"التيما", ("نيسان", "التيما", "Altima")),
]

MAKE_EN = {"تويوتا": "Toyota", "لكزس": "Lexus", "هيونداي": "Hyundai", "نيسان": "Nissan"}

# نوع القطعة -> (العربي، الإنجليزي، التصنيف الرئيسي)
PART_TYPES = [
    (r"مساعد", ("مساعد", "Shock Absorber", "نظام التعليق")),
    (r"مقص", ("مقص", "Control Arm", "نظام التعليق")),
    (r"عكس", ("عكس", "Tie Rod End", "نظام التوجيه")),
    (r"فحمات", ("فحمات فرامل", "Brake Pads", "الفرامل")),
    (r"هوبات|هوب", ("هوبات فرامل", "Brake Discs", "الفرامل")),
    (r"فلتر\s*زيت", ("فلتر زيت", "Oil Filter", "الفلاتر")),
    (r"فلتر\s*هواء", ("فلتر هواء", "Air Filter", "الفلاتر")),
    (r"فلتر\s*مكيف", ("فلتر مكيف", "Cabin Filter", "الفلاتر")),
    (r"بواجي|بوجي", ("بواجي", "Spark Plugs", "المحرك")),
    (r"قاعدة\s*مكينة", ("قاعدة مكينة", "Engine Mount", "المحرك")),
    (r"رديتر", ("رديتر", "Radiator", "التبريد")),
    (r"طرمبة\s*ماء", ("طرمبة ماء", "Water Pump", "التبريد")),
    (r"دينمو", ("دينمو", "Alternator", "الكهرباء")),
    (r"سلف", ("سلف", "Starter", "الكهرباء")),
    (r"كمبروسر", ("كمبروسر مكيف", "AC Compressor", "التكييف")),
]

POSITION = {"امامي": ("أمامي", "Front"), "خلفي": ("خلفي", "Rear")}
SIDE = [(r"يمين\s*\+\s*يسار", ("يمين ويسار", "Left & Right")),
        (r"يمين", ("يمين", "Right")), (r"يسار", ("يسار", "Left"))]

BRAND_ALIASES = [(r"K\s*\.?\s*Y\s*\.?\s*B\.?", "KYB")]


def norm_brand(s):
    s = (s or "").strip()
    for pat, b in BRAND_ALIASES:
        if re.search(pat, s, re.I):
            return b
    return s


def fix_year(y):
    y = int(y)
    if y < 100:
        y += 1900 if y > 30 else 2000
    return y


def parse_name(raw):
    name = re.sub(r"\s+", " ", raw).strip()
    name = re.sub(r"مساعد(?=[^\s])", "مساعد ", name)  # "مساعدخلفي"
    notes, work = [], name

    # السنوات (تُشال قبل استخراج الأرقام البديلة)
    y_from = y_to = None
    m = re.search(r"(\d{2,4})\s*-\s*(\d{2,4})?", work)
    if m:
        a = fix_year(m.group(1))
        b = fix_year(m.group(2)) if m.group(2) else None
        if b and b < a:
            a, b = b, a
        y_from, y_to = a, b
        work = work[:m.start()] + " " + work[m.end():]
    else:
        notes.append("سنوات غير موجودة")

    # الموديلات (ممكن أكثر من موديل: هايلكس + فورتشنر)
    fits = []
    for pat, info in MODELS:
        mm = re.search(pat, work, re.I)
        if mm:
            if info[0] == "لكزس" and not info[1] and any(f[0] == "لكزس" for f in fits):
                continue
            fits.append(info)
            work = work[:mm.start()] + " " + work[mm.end():]
    if not fits:
        notes.append("السيارة غير معروفة — أضفها للقاموس")
    if any(f[0] == "لكزس" and not f[1] for f in fits):
        notes.append("موديل لكزس غير محدد")

    trim = []
    if re.search(r"دبل", work):
        trim.append("غمارتين")
    if re.search(r"عايدي|عادي", work):
        trim.append("غمارة")

    part = next((info for pat, info in PART_TYPES if re.search(pat, work)), None)
    if not part:
        notes.append("نوع القطعة غير معروف")
    pos = next((v for k, v in POSITION.items() if k in work), None)
    side = next((v for pat, v in SIDE if re.search(pat, work)), None)
    if re.search(r"\sR\s*$", name) and not side:
        notes.append("حرف R في الاسم — تأكد من الجهة")

    # أرقام بديلة: أي رقم 4-7 خانات متبقي
    alts = re.findall(r"(?<![\d.])\d{4,7}(?![\d.])", work)
    origin = "ياباني" if "ياباني" in work else ""
    return dict(name=name, y_from=y_from, y_to=y_to, fits=fits, trim=trim, part=part,
                pos=pos, side=side, alts=alts, origin=origin, notes=notes)


def search_key(sku):
    sku = sku.strip()
    if re.fullmatch(r"\d{5}-\d{5}", sku):
        return sku, "OEM تويوتا"
    if re.fullmatch(r"J\d{6,7}", sku):
        return sku[1:], "رقم الماركة (بدون J)"
    return sku, "رقم الماركة"


def years_txt(p):
    if not p["y_from"]:
        return ""
    return f"{p['y_from']}-{p['y_to']}" if p["y_to"] else f"{p['y_from']} وأحدث"


def build(row_sku, raw_name, unit, price, brand):
    p = parse_name(raw_name)
    key, key_type = search_key(row_sku)
    yrs = years_txt(p)
    part_ar = p["part"][0] if p["part"] else ""
    part_en = p["part"][1] if p["part"] else ""
    pos_ar, pos_en = p["pos"] or ("", "")
    side_ar, side_en = p["side"] or ("", "")
    makes = list(dict.fromkeys(f[0] for f in p["fits"]))
    fit_ar = " / ".join(
        " ".join(x for x in [mk, " / ".join(f[1] for f in p["fits"] if f[0] == mk and f[1])] if x)
        for mk in makes)
    fit_en = " ".join(" ".join(x for x in (MAKE_EN.get(f[0], ""), f[2]) if x) for f in p["fits"])
    trim = " ".join(p["trim"])

    title = " ".join(x for x in [part_ar, pos_ar, side_ar, fit_ar, trim, yrs, brand] if x)
    meta_title = title if len(title) <= 60 else " ".join(x for x in [part_ar, pos_ar, side_ar, fit_ar, yrs] if x)[:60]
    alt_txt = "، ".join(p["alts"])
    meta_desc = (f"{part_ar} {brand} {pos_ar} {side_ar} لـ {fit_ar} {yrs}. رقم القطعة {key}"
                 + (f"، رقم بديل {alt_txt}" if alt_txt else "") + ". تأكد من التوافق برقم القطعة واطلب الآن.")
    meta_desc = re.sub(r"\s+", " ", meta_desc).replace(" .", ".")[:155]
    slug = "-".join(x for x in [brand, pos_en, side_en.replace(" & ", "-"), part_en, fit_en, yrs.replace(" وأحدث", "-plus"), key] if x)
    slug = re.sub(r"[^a-z0-9]+", "-", slug.lower()).strip("-")

    desc_lines = [
        f"{part_ar} {pos_ar} {side_ar} من {brand}" + (f" ({p['origin']})" if p["origin"] else ""),
        f"السيارات المتوافقة: {fit_ar} {trim} {yrs}",
        f"رقم القطعة: {key}",
    ]
    if alt_txt:
        desc_lines.append(f"أرقام بديلة: {alt_txt}")
    if row_sku != key:
        desc_lines.append(f"رقم الصنف: {row_sku}")
    desc = "\n".join(re.sub(r"\s+", " ", l).strip() for l in desc_lines)

    notes = list(p["notes"])
    try:
        price_f = float(price)
    except (TypeError, ValueError):
        price_f = 0.0
    if price_f <= 0:
        notes.append("السعر صفر — لا يُرفع")
    if side_ar == "يمين ويسار":
        notes.append("يمين+يسار: هل السعر للحبة أو للطقم؟")

    main_cat = p["part"][2] if p["part"] else "غير مصنف"
    sub_cat = f"{part_ar}ات" if part_ar == "مساعد" else part_ar
    group = "|".join([brand, part_ar, pos_ar, fit_ar, trim, yrs])

    return {
        "رقم الصنف (SKU)": row_sku,
        "رقم البحث في الكتالوج": key,
        "نوع الرقم": key_type,
        "ماركة القطعة": brand,
        "نوع القطعة": part_ar,
        "الموضع": pos_ar,
        "الجهة": side_ar,
        "ماركة السيارة": " / ".join(makes),
        "السيارات المتوافقة": " ".join(x for x in [fit_ar, trim] if x),
        "سنة من": p["y_from"],
        "سنة إلى": p["y_to"],
        "أرقام بديلة": alt_txt,
        "السعر": price_f,
        "الوحدة": unit,
        "اسم المنتج المحسّن": re.sub(r"\s+", " ", title),
        "Meta Title": re.sub(r"\s+", " ", meta_title),
        "Meta Description": meta_desc,
        "الرابط (Slug)": slug,
        "الوصف": desc,
        "التصنيف الرئيسي": main_cat,
        "التصنيف الفرعي": sub_cat,
        "تصنيف السيارة": " / ".join(makes),
        "_group": group,
        "منتج بخيارات (يمين/يسار)": "",
        "ملاحظات للمراجعة": " | ".join(notes),
        "الاسم الأصلي": raw_name.strip(),
    }


def read_report(path):
    ws = openpyxl.load_workbook(path, read_only=True, data_only=True).worksheets[0]
    brand, rows = "", []
    for r in ws.iter_rows(values_only=True):
        cells = [str(c).strip() if c is not None else "" for c in r]
        first = cells[0] if cells else ""
        if first.startswith("المجموعة"):
            brand = norm_brand(first.split(":", 1)[-1])
            continue
        if len(cells) < 5 or not cells[1] or cells[1] == "رقم الصنف" or not cells[2]:
            continue
        rows.append((cells[1], cells[2], cells[3], cells[4], brand))
    return rows


def main(src, dst):
    out = [build(*r) for r in read_report(src)]

    groups = defaultdict(list)
    for o in out:
        if o["الجهة"] in ("يمين", "يسار"):
            groups[o["_group"]].append(o)
    for g, items in groups.items():
        sides = {i["الجهة"] for i in items}
        if sides == {"يمين", "يسار"}:
            for i in items:
                i["منتج بخيارات (يمين/يسار)"] = "نعم — " + " + ".join(x["رقم الصنف (SKU)"] for x in items)

    cols = [k for k in out[0] if not k.startswith("_")]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "المنتجات"
    ws.sheet_view.rightToLeft = True
    ws.append(cols)
    warn = PatternFill("solid", fgColor="FFF2CC")
    bad = PatternFill("solid", fgColor="F8CBAD")
    for o in out:
        ws.append([o[c] for c in cols])
        if o["ملاحظات للمراجعة"]:
            fill = bad if "لا يُرفع" in o["ملاحظات للمراجعة"] else warn
            for cell in ws[ws.max_row]:
                cell.fill = fill
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3864")
    for i, c in enumerate(cols, 1):
        width = 60 if c in ("الوصف", "Meta Description") else 45 if "اسم" in c or c in ("Meta Title", "الرابط (Slug)") else 16
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"

    s = wb.create_sheet("ملخص")
    s.sheet_view.rightToLeft = True
    flagged = sum(1 for o in out if o["ملاحظات للمراجعة"])
    zero = sum(1 for o in out if o["السعر"] <= 0)
    paired = sum(1 for o in out if o["منتج بخيارات (يمين/يسار)"])
    for line in [
        ("إجمالي الأصناف", len(out)),
        ("جاهزة بدون ملاحظات", len(out) - flagged),
        ("تحتاج مراجعة", flagged),
        ("سعرها صفر", zero),
        ("أصناف تندمج كخيارات يمين/يسار", paired),
        ("عدد المنتجات في سلة بعد الدمج", len(out) - paired // 2),
    ]:
        s.append(line)
    s.column_dimensions["A"].width = 40
    wb.save(dst)
    print(f"{len(out)} صنف | {flagged} تحتاج مراجعة | {zero} سعر صفر | {paired} تندمج كخيارات -> {dst}")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
